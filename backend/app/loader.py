"""Read city workbooks (one .xlsx per city) into flat indicator records.

Expected layout, as produced by the Atlas city workbooks:
  * one sheet per theme (any sheet called "Summary" or "Change Log" is ignored)
  * a header row containing at least "Indicator" and "Value"; optional columns
    "Topic", "Rating", "Data basis", "Source & timeline", "Verification note"
  * a merged topic band row may precede each topic's indicators

The loader is forgiving: column order does not matter and extra columns are ignored.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

SKIP_SHEETS = {"summary", "change log", "readme", "notes"}
HEADER_ALIASES = {
    "indicator": "indicator",
    "topic": "topic",
    "value": "value",
    "rating": "rating",
    "data basis": "basis",
    "basis": "basis",
    "source & timeline": "source",
    "source": "source",
    "data source": "source",
    "verification note": "note",
    "note": "note",
}
RATING_LABELS = ["Exceptional", "High", "Moderate-High", "Moderate", "Low-Moderate", "Low"]
_RATING_NORMAL = {r.lower(): r for r in RATING_LABELS}
_RATING_NORMAL.update({"med-high": "Moderate-High", "medium": "Moderate", "moderate-low": "Low-Moderate"})


@dataclass
class Record:
    theme: str
    topic: str
    indicator: str
    value: str
    rating: str | None
    basis: str
    source: str
    note: str
    sheet: str
    row: int

    @property
    def key(self) -> str:
        return f"{self.theme}|{self.topic}|{self.indicator}"


@dataclass
class CityFile:
    city: str
    path: Path
    records: list[Record] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def city_from_filename(path: Path) -> str:
    stem = path.stem
    stem = re.sub(r"_?Urban_?Atlas$", "", stem, flags=re.I)
    return stem.replace("_", " ").strip() or path.stem


def normalise_rating(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s*\(.*\)", "", str(v)).strip().lower()
    return _RATING_NORMAL.get(s)


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_city_file(path: Path, city: str | None = None) -> CityFile:
    cf = CityFile(city=city or city_from_filename(path), path=path)
    wb = load_workbook(path, data_only=True, read_only=False)
    theme_sheets = 0
    for ws in wb.worksheets:
        if ws.title.strip().lower() in SKIP_SHEETS:
            continue
        header_row, cols = None, {}
        for r in range(1, min(ws.max_row, 15) + 1):
            found = {}
            for c in range(1, ws.max_column + 1):
                h = _clean(ws.cell(r, c).value).lower()
                if h in HEADER_ALIASES and HEADER_ALIASES[h] not in found:
                    found[HEADER_ALIASES[h]] = c
            if "indicator" in found and "value" in found:
                header_row, cols = r, found
                break
        if header_row is None:
            cf.warnings.append(f"Sheet '{ws.title}' skipped: no header row with 'Indicator' and 'Value'.")
            continue
        theme_sheets += 1
        theme = ws.title.strip()
        topic = ""
        for r in range(header_row + 1, ws.max_row + 1):
            name = _clean(ws.cell(r, cols["indicator"]).value)
            first = _clean(ws.cell(r, 1).value)
            if not name:
                # topic band: text in column A, nothing in the indicator column
                if first and not first.lower().startswith(("blue bold", "note", "values ")) and len(first) < 80:
                    topic = first
                continue
            if "topic" in cols and _clean(ws.cell(r, cols["topic"]).value):
                topic = _clean(ws.cell(r, cols["topic"]).value)
            raw_rating = ws.cell(r, cols["rating"]).value if "rating" in cols else None
            rating = normalise_rating(raw_rating)
            if raw_rating not in (None, "", "Not rated") and rating is None:
                cf.warnings.append(f"{theme} row {r}: rating '{raw_rating}' not recognised; treated as not rated.")
            cf.records.append(Record(
                theme=theme, topic=topic or theme, indicator=name,
                value=_clean(ws.cell(r, cols["value"]).value),
                rating=rating,
                basis=_clean(ws.cell(r, cols["basis"]).value) if "basis" in cols else "",
                source=_clean(ws.cell(r, cols["source"]).value) if "source" in cols else "",
                note=_clean(ws.cell(r, cols["note"]).value) if "note" in cols else "",
                sheet=ws.title, row=r,
            ))
    if theme_sheets == 0:
        cf.warnings.append("No theme sheets found; check the workbook layout.")
    return cf


def load_folder(folder: Path) -> list[CityFile]:
    files = sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith(("~$", ".")))
    return [read_city_file(p) for p in files]
