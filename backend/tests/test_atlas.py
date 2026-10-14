import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture()
def workspace(tmp_path, monkeypatch):
    data = tmp_path / "data"
    shutil.copytree(ROOT / "data" / "cities", data / "cities")
    shutil.copytree(ROOT / "config", tmp_path / "config")
    return data, tmp_path / "config"


def set_value(path, sheet, indicator_prefix, value):
    wb = load_workbook(path)
    ws = wb[sheet]
    for row in ws.iter_rows():
        if row[2].value and str(row[2].value).startswith(indicator_prefix):
            row[3].value = value
    wb.save(path)


def test_loader_reads_every_city(workspace):
    from app.loader import load_folder
    data, _ = workspace
    files = load_folder(data / "cities")
    assert len(files) == 8
    for f in files:
        assert len(f.records) == 202, f.city
        assert not f.warnings, f.warnings


def test_rank_label_matches_eight_city_bands():
    from app.engine import rank_label
    assert [rank_label(r, 8) for r in range(1, 9)] == [
        "Exceptional", "High", "Moderate-High", "Moderate", "Moderate", "Low-Moderate", "Low", "Low"]


def test_parse_number():
    from app.engine import parse_number
    assert parse_number("₹4,93,024") == 493024
    assert parse_number("1.60x (TomTom 2025: 60.2%)") == 1.6
    assert parse_number("N/A (not an SCM city)") is None
    assert parse_number("—") is None


def test_edit_recomputes_ratings_and_logs_history(workspace):
    from app.store import Store
    data, config = workspace
    store = Store(data, config)
    ds = store.get()
    piped = next(i for i in ds["indicators"] if i["name"].startswith("Households Connected to Piped"))
    assert piped["ratings"]["Patna"] == "Low"
    set_value(data / "cities" / "Patna_Urban_Atlas.xlsx", "Urban Infrastructure", "Households Connected to Piped", "99%")
    ds = store.get()
    piped = next(i for i in ds["indicators"] if i["name"].startswith("Households Connected to Piped"))
    assert piped["values"]["Patna"] == "99%"
    assert piped["ratings"]["Patna"] == "Exceptional"
    changes = [e for e in ds["history"] if e["type"] == "value"]
    assert changes and changes[0]["city"] == "Patna" and changes[0]["old"] == "48%"


@pytest.fixture()
def client(workspace, monkeypatch):
    data, config = workspace
    monkeypatch.setenv("ATLAS_DATA_DIR", str(data))
    monkeypatch.setenv("ATLAS_CONFIG_DIR", str(config))
    monkeypatch.setenv("ATLAS_ADMIN_USER", "boss")
    monkeypatch.setenv("ATLAS_ADMIN_PASSWORD", "admin-pass-123")
    monkeypatch.delenv("ATLAS_PUBLIC_READ", raising=False)
    import importlib
    import app.main as main
    importlib.reload(main)
    from fastapi.testclient import TestClient
    with TestClient(main.app) as c:  # runs start-up, which creates the first admin
        main.auth.add_user("reader", "viewer-pass-123", "viewer")
        yield c, data

H = {"X-Requested-With": "atlas"}


def login(c, user, pw):
    r = c.post("/api/auth/login", json={"username": user, "password": pw})
    assert r.status_code == 200, r.text
    return r.json()["user"]


def test_reading_needs_sign_in(client):
    c, _ = client
    assert c.get("/api/dashboard").status_code == 401
    assert c.post("/api/auth/login", json={"username": "boss", "password": "wrong-password"}).status_code == 401
    assert login(c, "reader", "viewer-pass-123")["role"] == "viewer"
    d = c.get("/api/dashboard").json()
    assert d["cities"][0] == "Ahmedabad" and d["can_write"] is False
    assert c.get("/api/cities/patna").json()["readiness"]["verdict"] == "hold"
    assert c.get("/api/export.csv").text.count("\n") == 8 * 202 + 1
    c.post("/api/auth/logout")
    assert c.get("/api/dashboard").status_code == 401


def test_viewer_cannot_write(client):
    c, data = client
    login(c, "reader", "viewer-pass-123")
    with (data / "cities" / "Kanpur_Urban_Atlas.xlsx").open("rb") as fh:
        assert c.post("/api/admin/upload", files={"file": ("Kanpur_Urban_Atlas.xlsx", fh)}, headers=H).status_code == 403
    assert c.post("/api/admin/reload", headers=H).status_code == 403
    assert c.get("/api/admin/users", headers=H).status_code == 403


def test_admin_upload_and_user_management(client):
    c, data = client
    assert login(c, "boss", "admin-pass-123")["role"] == "admin"
    assert c.get("/api/dashboard").json()["can_write"] is True
    up = data.parent / "Kanpur_Urban_Atlas.xlsx"
    shutil.copy2(data / "cities" / "Kanpur_Urban_Atlas.xlsx", up)
    set_value(up, "Urban Infrastructure", "Households Connected to Piped", "60%")
    with up.open("rb") as fh:  # missing anti-forgery header
        assert c.post("/api/admin/upload", files={"file": (up.name, fh)}).status_code == 400
    with up.open("rb") as fh:
        r = c.post("/api/admin/upload", files={"file": (up.name, fh)}, headers=H)
    assert r.status_code == 200, r.text
    assert r.json()["values_changed"] == 1
    assert "upload by boss" in c.get("/api/history").json()[0]["reason"]

    assert c.post("/api/admin/users", json={"username": "kothari", "password": "pw-123456", "role": "viewer"}, headers=H).status_code == 200
    assert {u["username"] for u in c.get("/api/admin/users", headers=H).json()} == {"boss", "reader", "kothari"}
    assert c.patch("/api/admin/users/kothari", json={"role": "admin"}, headers=H).json()["role"] == "admin"
    assert c.delete("/api/admin/users/boss", headers=H).status_code == 400  # not your own account
    assert c.delete("/api/admin/users/kothari", headers=H).status_code == 200


def test_password_reset_ends_other_sessions(client):
    c, _ = client
    login(c, "reader", "viewer-pass-123")
    assert c.get("/api/dashboard").status_code == 200
    import app.main as main
    main.auth.update_user("reader", password="brand-new-pass")
    assert c.get("/api/dashboard").status_code == 401


def test_passwords_are_hashed(client):
    _, data = client
    text = (data / "users.json").read_text()
    assert "viewer-pass-123" not in text and "pbkdf2_sha256$" in text
